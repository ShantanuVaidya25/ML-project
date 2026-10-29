"""Renders the report (report_content.py) to Project_Report.docx, editable in Word / Google Docs.

    python make_report_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from report_content import build

ROOT = Path(__file__).parent
OUT = ROOT / "Project_Report.docx"
BLUE, GREY = RGBColor(0x1F, 0x5F, 0xBF), RGBColor(0x5A, 0x5A, 0x5A)
TEXT_W = Mm(170)  # A4 width minus 20 mm margins


def runs(par, text, size=None, italic=False):
    """Add text with <b>/<i> markup as runs."""
    for tok in re.split(r"(</?[bi]>)", text):
        if tok in ("<b>", "</b>", "<i>", "</i>"):
            runs.state[tok[-2]] = not tok.startswith("</")
            continue
        if tok:
            r = par.add_run(tok)
            r.bold, r.italic = runs.state["b"], runs.state["i"] or italic
            if size:
                r.font.size = Pt(size)
    runs.state = {"b": False, "i": False}


runs.state = {"b": False, "i": False}


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(p, text)
    return p


def table(doc, rows, widths, center=False):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER if center else WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = int(TEXT_W * widths[j])
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            runs(p, f"<b>{val}</b>" if i == 0 else str(val), size=9)
            if i == 0:
                shade(c, "EEF3FB")
    # repeat header row on page breaks
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    tr_pr.append(hdr)
    doc.add_paragraph()
    return t


def page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Real-Time PPE Detection - Mini Project Report    ")
    r.font.size, r.font.color.rgb = Pt(8), GREY
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        run = p.add_run()
        run.font.size = Pt(8)
        if tag:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), tag)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._r.append(instr)


def main():
    blocks, have = build(ROOT)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Mm(20))
    sec.top_margin = sec.bottom_margin = Mm(18)
    sec.different_first_page_header_footer = True
    page_number_footer(sec)

    st = doc.styles
    st["Normal"].font.name, st["Normal"].font.size = "Times New Roman", Pt(11.5)
    st["Normal"].paragraph_format.space_after = Pt(6)
    st["Normal"].paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        st[name].font.name, st[name].font.size, st[name].font.bold = "Arial", Pt(size), True
        st[name].font.color.rgb = BLUE if name == "Heading 1" else RGBColor(0, 0, 0)
    st["Caption"].font.size, st["Caption"].font.color.rgb, st["Caption"].font.italic = Pt(9), GREY, False

    nfig = ntab = 0
    for b in blocks:
        kind = b[0]
        if kind == "cover":
            info = b[1]
            for _ in range(4):
                doc.add_paragraph()
            for part in info["title"].split(": ", 1):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(part)
                r.bold, r.font.size, r.font.name = True, Pt(22), "Arial"
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(info["subtitle"])
            r.font.size, r.font.color.rgb = Pt(13), GREY
            doc.add_paragraph()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run("Submitted by").bold = True
            table(doc, [["Name", "Roll No."]] + [list(m) for m in info["team"]], [0.45, 0.2], center=True)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run("Under the guidance of").bold = True
            p.add_run().add_break()
            p.add_run(info["guide"])
            doc.add_paragraph()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            runs(p, f"<b>Abstract.</b> {info['abstract']}")
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "h1":
            doc.add_heading(b[1], level=1)
        elif kind == "h2":
            doc.add_heading(b[1], level=2)
        elif kind == "p":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            runs(p, b[1])
        elif kind == "todo":
            p = doc.add_paragraph()
            r = p.add_run(f"[{b[1]}]")
            r.italic, r.font.color.rgb = True, RGBColor(0xB0, 0x70, 0x00)
        elif kind == "bullets":
            for it in b[1]:
                p = doc.add_paragraph(style="List Number" if b[2] else "List Bullet")
                runs(p, it)
        elif kind == "table":
            ntab += 1
            caption(doc, f"Table {ntab}. {b[1]}").paragraph_format.keep_with_next = True
            table(doc, b[2], b[3])
        elif kind == "figure":
            path = ROOT / b[1]
            if not path.exists():
                doc.add_paragraph(f"[Figure missing: {b[1]}]")
                continue
            nfig += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(path), width=int(TEXT_W * b[3]))
            caption(doc, f"Figure {nfig}. {b[2]}")
        elif kind == "refs":
            for i, r in enumerate(b[1], 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Mm(8)
                p.paragraph_format.first_line_indent = Mm(-8)
                runs(p, f"[{i}] {r}", size=10)

    doc.core_properties.title = "Real-Time PPE Detection - Mini Project Report"
    doc.core_properties.author = "Shantanu Vaidya, Vineet Shah, Mohammad Ejaz Mehendi"
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
